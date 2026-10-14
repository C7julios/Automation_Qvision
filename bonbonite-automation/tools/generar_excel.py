# -*- coding: utf-8 -*-
"""Genera PTP_Carlos_Sanchez_13102026.xlsx a partir de los features y del reporte JSON de Serenity."""
import glob
import json
import os
import shutil
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from PIL import Image

# Carpeta que contiene bonbonite-automation/ (por defecto, dos niveles arriba de este script)
RAIZ = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PROYECTO = os.path.join(RAIZ, "bonbonite-automation")
REPORTE = os.path.join(PROYECTO, "target", "site", "serenity")
EVIDENCIAS = os.path.join(RAIZ, "evidencias")
SALIDA = os.path.join(RAIZ, "PTP_Carlos_Sanchez_13102026.xlsx")

# ---------- estilos ----------
FUENTE = "Arial"
NEGRO, AMARILLO, GRIS, GRIS_CLARO, BLANCO = "1F1F1F", "F6CB5B", "595959", "F2F2F2", "FFFFFF"
VERDE, ROJO = "D9EAD3", "F4CCCC"
borde = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
f_titulo = Font(name=FUENTE, size=18, bold=True, color=NEGRO)
f_subtitulo = Font(name=FUENTE, size=11, italic=True, color=GRIS)
f_seccion = Font(name=FUENTE, size=12, bold=True, color=NEGRO)
f_header = Font(name=FUENTE, size=10, bold=True, color=BLANCO)
f_normal = Font(name=FUENTE, size=10, color=NEGRO)
f_bold = Font(name=FUENTE, size=10, bold=True, color=NEGRO)
f_code = Font(name="Consolas", size=10, color=NEGRO)
fill_header = PatternFill("solid", fgColor=NEGRO)
fill_seccion = PatternFill("solid", fgColor=AMARILLO)
fill_label = PatternFill("solid", fgColor=GRIS_CLARO)
fill_ok = PatternFill("solid", fgColor=VERDE)
fill_ko = PatternFill("solid", fgColor=ROJO)
wrap = Alignment(wrap_text=True, vertical="top")
wrap_c = Alignment(wrap_text=True, vertical="top", horizontal="center")


def ancho(ws, anchos):
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def titulo(ws, texto, sub, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(1, 1, texto).font = f_titulo
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(2, 1, sub).font = f_subtitulo
    ws.sheet_view.showGridLines = False
    return 4


def seccion(ws, fila, texto, ncols):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    c = ws.cell(fila, 1, texto)
    c.font, c.fill = f_seccion, fill_seccion
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[fila].height = 20
    return fila + 1


def alto(textos_por_ancho):
    """Estimación simple de alto de fila según el texto más largo y el ancho de su columna."""
    lineas = 1
    for texto, w in textos_por_ancho:
        if texto is None:
            continue
        n = 0
        for parte in str(texto).split("\n"):
            n += max(1, -(-len(parte) // max(1, int(w * 1.1))))
        lineas = max(lineas, n)
    return max(15, 13.5 * lineas)


def tabla(ws, fila, encabezados, filas, anchos_col, merges=None, estado_col=None):
    """Escribe una tabla con encabezado negro. merges: {indice_col: n_cols_a_unir}."""
    merges = merges or {}
    col = 1
    posiciones = []
    for i, h in enumerate(encabezados):
        span = merges.get(i, 1)
        posiciones.append((col, span))
        if span > 1:
            ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + span - 1)
        for k in range(span):
            ws.cell(fila, col + k).fill = fill_header
            ws.cell(fila, col + k).border = borde
        c = ws.cell(fila, col, h)
        c.font, c.alignment = f_header, Alignment(wrap_text=True, vertical="center", horizontal="center")
        col += span
    ws.row_dimensions[fila].height = 28
    fila += 1
    for datos in filas:
        textos = []
        for i, v in enumerate(datos):
            c0, span = posiciones[i]
            if span > 1:
                ws.merge_cells(start_row=fila, start_column=c0, end_row=fila, end_column=c0 + span - 1)
            for k in range(span):
                ws.cell(fila, c0 + k).border = borde
            c = ws.cell(fila, c0, v)
            c.font = f_bold if i == 0 else f_normal
            c.alignment = wrap_c if (estado_col is not None and i == estado_col) else wrap
            if estado_col is not None and i == estado_col:
                c.fill = fill_ok if str(v).lower().startswith("pas") else fill_ko
            w = sum(anchos_col[c0 - 1 + k] for k in range(span))
            textos.append((v, w))
        ws.row_dimensions[fila].height = alto(textos)
        fila += 1
    return fila + 1


def ficha(ws, fila, pares, ncols, anchos_col):
    """Pares etiqueta/valor; la etiqueta en col A y el valor unido de B hasta ncols."""
    w_val = sum(anchos_col[1:ncols])
    for etiqueta, valor in pares:
        a = ws.cell(fila, 1, etiqueta)
        a.font, a.fill, a.alignment, a.border = f_bold, fill_label, wrap, borde
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=ncols)
        for k in range(2, ncols + 1):
            ws.cell(fila, k).border = borde
        b = ws.cell(fila, 2, valor)
        b.font, b.alignment = f_normal, wrap
        ws.row_dimensions[fila].height = alto([(valor, w_val), (etiqueta, anchos_col[0])])
        fila += 1
    return fila + 1


def bloque_codigo(ws, fila, texto, ncols):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    c = ws.cell(fila, 1, texto)
    c.font, c.alignment = f_code, Alignment(wrap_text=True, vertical="top")
    c.fill = PatternFill("solid", fgColor="FAFAFA")
    for k in range(1, ncols + 1):
        ws.cell(fila, k).border = borde
    ws.row_dimensions[fila].height = 14 * (texto.count("\n") + 1) + 6
    return fila + 2


# ---------- datos de la ejecución ----------
def cargar_resultados():
    resultados = {}
    for f in glob.glob(os.path.join(REPORTE, "*.json")):
        try:
            d = json.load(open(f, encoding="utf8"))
        except Exception:
            continue
        if "testSteps" not in d or "title" not in d:
            continue
        pasos = d["testSteps"]
        # Los escenarios con Ejemplos traen un nivel adicional por fila de datos
        if len(pasos) == 1 and pasos[0]["description"].startswith("1:"):
            pasos = pasos[0]["children"]
        resultados[d["userStory"]["storyName"] if "userStory" in d else d["title"]] = {
            "titulo": d["title"],
            "resultado": d["result"],
            "duracion": d.get("duration", 0) / 1000.0,
            "inicio": d.get("startTime", "")[:19].replace("T", " "),
            "pasos": [
                {
                    "descripcion": p["description"],
                    "resultado": p["result"],
                    "duracion": p.get("duration", 0) / 1000.0,
                    "acciones": [h["description"] for h in p.get("children", [])],
                    "captura": (p.get("screenshots") or [{}])[-1].get("screenshot"),
                }
                for p in pasos
            ],
        }
    return resultados


def leer_feature(nombre):
    return open(os.path.join(PROYECTO, "src", "test", "resources", "features", nombre), encoding="utf8").read().strip()


def copiar_evidencia(nombre_png, destino):
    origen = os.path.join(REPORTE, nombre_png)
    os.makedirs(EVIDENCIAS, exist_ok=True)
    ruta = os.path.join(EVIDENCIAS, destino)
    shutil.copyfile(origen, ruta)
    return ruta


def insertar_imagen(ws, ruta, celda, ancho_px=760):
    with Image.open(ruta) as im:
        w, h = im.size
    img = XLImage(ruta)
    img.width, img.height = ancho_px, int(h * ancho_px / w)
    ws.add_image(img, celda)
    return img.height


ESTADO = {"SUCCESS": "Pasó", "FAILURE": "Falló", "ERROR": "Error", "SKIPPED": "Omitido", "PENDING": "Pendiente"}

resultados = cargar_resultados()
por_feature = {}
for r in resultados.values():
    t = r["titulo"].lower()
    clave = "registro" if "registro" in t else "modificar" if "modificar" in t else "compra"
    por_feature[clave] = r
faltan = {"registro", "modificar", "compra"} - set(por_feature)
if faltan:
    sys.exit("No se encontraron resultados de: " + ", ".join(faltan))

fecha_ejecucion = por_feature["registro"]["inicio"][:10]

wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)

# =====================================================================
# Hoja 1: Portada
# =====================================================================
ws = wb.active
ws.title = "Portada"
A = [30, 44.44140625, 22, 22, 22]
ancho(ws, A)
f = titulo(ws, "Plan de Pruebas Automatizadas – Sitio web Bon-bonite v1",
           "Proceso de pruebas desde cero · Automatización de 3 escenarios con Serenity BDD + Screenplay", 5)
f = seccion(ws, f, "Información general", 5)
f = ficha(ws, f, [
    ("Cliente", "Bon-bonite – Dirección de Mercadeo y Ventas"),
    ("Sistema bajo prueba", "Sitio web de comercio electrónico https://www.bon-bonite.com/ (versión 1, WooCommerce/WordPress)"),
    ("Documento", "Alcance, estrategia de pruebas y solución de tres escenarios automatizados"),
    ("Versión del documento", "1.0"),
    ("Fecha de ejecución", fecha_ejecucion),
    ("Proyecto de automatización", "bonbonite-automation/ (Maven · Java 17+ · Serenity BDD 4.2.1 · Cucumber 7 · Selenium 4)"),
    ("Reporte de ejecución", r"bonbonite-automation\target\site\serenity\index.html  ·  capturas en la carpeta evidencias\ "),
    ("Repositorio", "https://github.com/C7julios/Automation_Qvision"),
], 5, A)

f = seccion(ws, f, "Resumen de la ejecución", 5)
fila_ini = f + 1
f = tabla(ws, f, ["ID", "Escenario", "Hoja", "Estado", "Duración (s)"], [
    ["ESC-01", "Registro exitoso de usuario", "Esc1 - Registro", None, None],
    ["ESC-02", "Modificación de datos del usuario registrado", "Esc2 - Modificar datos", None, None],
    ["ESC-03", "Compra de producto (categoría Zapatos)", "Esc3 - Compra", None, None],
], A)
# Estado y duración enlazados a cada hoja de escenario (celdas B12 y B13 de la ficha)
for i, hoja in enumerate(["Esc1 - Registro", "Esc2 - Modificar datos", "Esc3 - Compra"]):
    r = fila_ini + i
    ws.cell(r, 4, f"='{hoja}'!B12").font = Font(name=FUENTE, size=10, color="008000")
    ws.cell(r, 5, f"='{hoja}'!B13").font = Font(name=FUENTE, size=10, color="008000")
    ws.cell(r, 4).alignment = wrap_c
    ws.cell(r, 5).number_format = "0.0"
r = fila_ini + 3
for c, (et, form) in enumerate([("Escenarios ejecutados", f"=COUNTA(A{fila_ini}:A{fila_ini + 2})"),
                                ("Escenarios exitosos", f'=COUNTIF(D{fila_ini}:D{fila_ini + 2},"Pasó")'),
                                ("% de éxito", f"=IF(B{r}=0,0,B{r + 1}/B{r})")]):
    ws.cell(r + c, 1, et).font = f_bold
    ws.cell(r + c, 1).fill = fill_label
    cel = ws.cell(r + c, 2, form)
    cel.font = f_bold
    ws.cell(r + c, 1).border = cel.border = borde
cel.number_format = "0%"
f = r + 4
ws.cell(f, 1, "Nota: los valores del resumen son fórmulas enlazadas a cada hoja de escenario; Excel los calcula al abrir.").font = f_subtitulo
f += 2
f = seccion(ws, f, "Contenido del documento", 5)
f = tabla(ws, f, ["Hoja", "Contenido"], [
    ["Alcance y Estrategia", "a) Objetivo, alcance, módulos, tipos y niveles de prueba, enfoque, herramientas, ambiente, datos, criterios, riesgos, roles, cronograma y entregables."],
    ["Esc1 - Registro", "b) Solución del escenario 1: ficha, Gherkin, pasos automatizados con resultados, componentes Screenplay y evidencias."],
    ["Esc2 - Modificar datos", "b) Solución del escenario 2 (misma estructura)."],
    ["Esc3 - Compra", "b) Solución del escenario 3 (misma estructura)."],
    ["Hallazgos", "Observaciones y posibles defectos detectados durante el diseño y la ejecución."],
], A, merges={1: 4})

# =====================================================================
# Hoja 2: Alcance y Estrategia
# =====================================================================
ws = wb.create_sheet("Alcance y Estrategia")
A = [26, 30, 30, 30, 30]
ancho(ws, A)
f = titulo(ws, "a) Alcance y estrategia de pruebas",
           "Proceso de pruebas desde cero para la versión 1 del sitio web Bon-bonite", 5)

f = seccion(ws, f, "1. Objetivo", 5)
f = ficha(ws, f, [
    ("Objetivo general", "Verificar que los flujos críticos del sitio web v1 funcionen correctamente de punta a punta, "
                         "para mejorar su calidad antes de construir la versión 2, y dejar una base de regresión "
                         "automatizada que se pueda ejecutar en cada cambio del sitio."),
    ("Objetivos específicos", "• Confirmar que un cliente potencial se puede registrar con éxito y queda autenticado.\n"
                              "• Confirmar que el usuario registrado puede iniciar sesión con su cédula y modificar sus datos, y que los cambios persisten.\n"
                              "• Confirmar que un cliente puede elegir un producto disponible, agregarlo al carrito y llegar al pago con el resumen correcto.\n"
                              "• Dejar documentado el proceso (alcance, estrategia, criterios, riesgos) para escalar la cobertura a todos los módulos."),
], 5, A)

f = seccion(ws, f, "2. Alcance", 5)
f = tabla(ws, f, ["Dentro del alcance", "Fuera del alcance"], [
    ["• Pruebas funcionales E2E sobre la interfaz web (escritorio, Chrome).\n"
     "• Registro de usuario (Mi cuenta / Account).\n"
     "• Inicio y cierre de sesión con cédula y contraseña.\n"
     "• Modificación de datos personales (pestaña Datos de Mi cuenta).\n"
     "• Navegación por categoría, selección de talla, carrito y checkout hasta el resumen del pedido.\n"
     "• Matriz de cobertura de los módulos Zapatos, Bolsos, Cinturones, Accesorios, Outlet, PQRS, Mi cuenta, Bonos de regalo y Account.",
     "• Confirmación real del pedido y pago (pasarela): se valida hasta el botón \"Registrar Orden\" para no crear pedidos en producción.\n"
     "• Pruebas de rendimiento, carga y estrés.\n"
     "• Pruebas de seguridad (OWASP, pentesting).\n"
     "• Integraciones internas de las áreas (Costos, Producción, Facturación, Contabilidad), que no existen en la v1.\n"
     "• Navegadores distintos de Chrome y dispositivos móviles (se proponen para la siguiente fase).\n"
     "• Recepción de correos (confirmación de registro / pedido)."],
], A, merges={0: 2, 1: 3})

f = seccion(ws, f, "3. Matriz de cobertura por módulo", 5)
f = tabla(ws, f, ["Módulo", "Escenarios candidatos", "Prioridad", "Automatizado en esta fase", "Escenario"], [
    ["Account (registro)", "Registro exitoso; registro con cédula/correo duplicado; contraseñas no coinciden; sin autorizar tratamiento de datos", "Alta", "Sí", "ESC-01"],
    ["Mi cuenta", "Login con cédula; logout; modificar datos personales; cambiar contraseña; direcciones de facturación/envío; recuperar contraseña", "Alta", "Sí (login, logout, datos)", "ESC-01, ESC-02, ESC-03"],
    ["Zapatos", "Listar productos; filtrar por subcategoría/talla/color; seleccionar talla con stock; agregar al carrito; compra", "Alta", "Sí", "ESC-03"],
    ["Bolsos", "Listar, filtrar, ver detalle, agregar al carrito (producto sin talla)", "Media", "No (siguiente fase)", "—"],
    ["Cinturones", "Listar, ver detalle, seleccionar talla, agregar al carrito", "Media", "No (siguiente fase)", "—"],
    ["Accesorios", "Listar, ver detalle, agregar al carrito", "Media", "No (siguiente fase)", "—"],
    ["Outlet", "Precio con descuento visible; precio tachado; total del carrito con descuento", "Alta", "No (siguiente fase)", "—"],
    ["Bonos de regalo", "Seleccionar denominación (50.000 a 300.000); agregar al carrito; checkout", "Media", "No (siguiente fase)", "—"],
    ["PQRS", "Radicar PQRS desde el sitio y desde Mi cuenta; validaciones de campos obligatorios", "Media", "No (siguiente fase)", "—"],
    ["Carrito / Checkout", "Cantidades, eliminar ítem, totales, IVA, departamento/ciudad, términos, registrar orden (en staging)", "Alta", "Parcial (hasta el resumen del pedido)", "ESC-03"],
], A)

f = seccion(ws, f, "4. Estrategia de pruebas", 5)
f = ficha(ws, f, [
    ("Tipos de prueba", "• Funcionales de caja negra E2E (UI).\n• Regresión automatizada de los flujos críticos.\n"
                        "• Exploratorias guiadas por sesiones para los módulos sin automatizar.\n• Pruebas de humo (smoke) en cada despliegue con los 3 escenarios."),
    ("Niveles", "Pruebas de sistema y de aceptación (E2E sobre el sitio desplegado). Se recomienda que la v2 añada pruebas unitarias y de API (REST de WooCommerce)."),
    ("Técnicas de diseño", "Partición de equivalencia y valores límite (datos de registro), tablas de decisión (validaciones del formulario), "
                           "transición de estados (anónimo → registrado → autenticado → carrito → checkout) y casos de uso por historia de usuario."),
    ("Enfoque de automatización", "BDD con Gherkin en español (lenguaje de negocio legible para Mercadeo y Ventas) + patrón Screenplay "
                                  "(actores, tareas, interacciones, preguntas). Screenplay separa el \"qué\" del \"cómo\": los cambios de la UI "
                                  "solo afectan las clases de userinterfaces/, lo que reduce el mantenimiento."),
    ("Criterio de selección de escenarios", "Se automatizaron los flujos de mayor riesgo e impacto de negocio: sin registro no hay cliente, sin "
                                            "datos correctos no hay entrega y sin compra no hay venta. Cubren lo que pide el caso: registro exitoso, "
                                            "modificación de datos y compra de producto."),
    ("Manejo de datos de prueba", "Datos dinámicos generados en cada ejecución (Datafaker): cédula única de 10 dígitos que empieza por 9 y "
                                  "correo qa.auto.<timestamp>@carlosmail.com para identificar las cuentas de la automatización. "
                                  "El producto no está quemado: se elige el primero de la categoría con talla en inventario."),
    ("Estabilidad", "Esperas explícitas (WaitUntil / WebDriverWait), sin XPath absolutos, localizadores por id/name/clases funcionales, "
                    "manejo del banner de cookies y de los selects con select2, escenarios independientes entre sí (cada uno crea su usuario)."),
], 5, A)

f = seccion(ws, f, "5. Herramientas y ambiente", 5)
f = tabla(ws, f, ["Elemento", "Detalle"], [
    ["Lenguaje / build", "Java 17+ (ejecutado con JDK 22) · Maven 3.9"],
    ["Framework", "Serenity BDD 4.2.1 (Screenplay) · Cucumber 7.18 · JUnit Platform 1.10"],
    ["Navegador / driver", "Google Chrome 155 · Selenium 4.23 (Selenium Manager descarga el driver)"],
    ["Reportes", "Serenity Living Documentation (target/site/serenity/index.html) con captura por acción"],
    ["Ambiente", "Producción: https://www.bon-bonite.com (no existe un ambiente de pruebas publicado). Se recomienda habilitar uno de staging."],
    ["Ejecución", "mvn clean verify  ·  por escenario: mvn clean verify -Dtags=\"@Compra\"  ·  sin interfaz: -Dheadless.mode=true"],
], A, merges={1: 4})

f = seccion(ws, f, "6. Criterios", 5)
f = tabla(ws, f, ["Tipo", "Criterio"], [
    ["Entrada", "Sitio disponible; ambiente de ejecución con JDK, Maven y Chrome; registro de usuarios habilitado; productos con inventario en la categoría Zapatos."],
    ["Salida", "100% de los escenarios automatizados ejecutados; 100% de los escenarios críticos aprobados o con su defecto reportado; reporte Serenity y evidencias publicados."],
    ["Suspensión", "Sitio caído o lento (>30 s por página); captcha o bloqueo de bots en registro/login; ausencia de inventario; cambios de UI que rompen más del 50% de los pasos."],
    ["Reanudación", "Corrección del bloqueo y ejecución de la prueba de humo con resultado exitoso."],
], A, merges={1: 4})

f = seccion(ws, f, "7. Riesgos y mitigación", 5)
f = tabla(ws, f, ["Riesgo", "Impacto", "Mitigación"], [
    ["Pruebas sobre producción (usuarios reales creados)", "Datos de prueba en la base de clientes; métricas de marketing alteradas", "Prefijo qa.auto en correos para depurarlos; no se confirma ningún pedido; solicitar ambiente de staging."],
    ["Inventario variable de productos y tallas", "El escenario de compra falla si el producto elegido se agota", "Selección dinámica del primer producto con talla en stock leyendo el JSON de variaciones."],
    ["Cambios frecuentes de la UI (tema propio bb-2024)", "Localizadores rotos", "Patrón Screenplay con localizadores centralizados en userinterfaces/."],
    ["Elementos de terceros (cookies, WhatsApp, chat)", "Clicks interceptados", "Interacción AceptarCookies y click por JavaScript en elementos cubiertos."],
    ["Captcha o protección anti-bots futura", "Bloqueo de la automatización", "Lista blanca de IP del runner o llave de prueba del captcha en staging."],
], A, merges={1: 2, 2: 2})

f = seccion(ws, f, "8. Roles, cronograma y entregables", 5)
f = tabla(ws, f, ["Rol", "Responsabilidad"], [
    ["Directora de Mercadeo y Ventas", "Patrocinadora; aprueba alcance, prioriza escenarios y recibe resultados."],
    ["Analista QA / Automatizador (consultor)", "Diseña la estrategia, automatiza, ejecuta, reporta resultados y defectos."],
    ["Desarrollador web", "Corrige defectos y provee ambiente de staging y datos de prueba."],
    ["Directores de área", "Validan los flujos que impactan su área (Facturación, Producción y logística, etc.)."],
], A, merges={1: 4})
f = tabla(ws, f, ["Fase", "Actividades", "Duración estimada"], [
    ["1. Análisis", "Entendimiento del negocio, recorrido del sitio, identificación de módulos y riesgos", "2 días"],
    ["2. Diseño", "Alcance, estrategia, matriz de cobertura, escenarios Gherkin", "2 días"],
    ["3. Automatización", "Proyecto Serenity + Screenplay, 3 escenarios, estabilización", "5 días"],
    ["4. Ejecución y reporte", "Ejecución, análisis de resultados, registro de hallazgos", "1 día"],
    ["5. Cierre", "Entrega de documento, código, reporte y recomendaciones para la v2", "1 día"],
], A, merges={1: 3})
f = tabla(ws, f, ["Entregable", "Ubicación"], [
    ["Documento de pruebas (este archivo)", "PTP_Carlos_Sanchez_13102026.xlsx"],
    ["Código de automatización", "bonbonite-automation/ (README.md con instrucciones)"],
    ["Reporte de ejecución Serenity", "bonbonite-automation/target/site/serenity/index.html (se genera al ejecutar)"],
    ["Evidencias", "evidencias/ (capturas de cada escenario)"],
], A, merges={1: 4})

# =====================================================================
# Hojas de escenarios
# =====================================================================
ESCENARIOS = [
    {
        "hoja": "Esc1 - Registro", "clave": "registro", "id": "ESC-01", "feature": "registro.feature",
        "nombre": "Registro exitoso de un nuevo usuario",
        "modulo": "Account / Mi cuenta",
        "historia": "Como cliente potencial de Bon-bonite quiero registrarme en el sitio web para poder iniciar sesión y comprar productos.",
        "prioridad": "Alta", "tipo": "Funcional E2E – camino feliz",
        "precondiciones": "• Sitio disponible.\n• Usuario no autenticado.\n• Cédula y correo no registrados previamente.",
        "datos": "Generados en cada ejecución (GeneradorDatos):\n• Nombres: Laura · Apellidos: aleatorio\n• Cédula: 10 dígitos únicos que empiezan por 9\n"
                 "• Correo: qa.auto.<timestamp>@carlosmail.com\n• Contraseña: QaAuto#<timestamp>\n• Autorización de tratamiento de datos: marcada",
        "esperado": "El usuario queda registrado y autenticado: ve el panel de Mi cuenta con el saludo \"Hola, Laura.\" y el menú de la cuenta muestra \"Cerrar Sesión\".",
        "esperados_pasos": [
            "Se muestra la página Mi cuenta con los formularios de ingreso y registro.",
            "El formulario se envía sin errores y se redirige al panel de la cuenta.",
            "La URL contiene /mi-cuenta y el saludo contiene el nombre del usuario.",
            "El menú de la cuenta se despliega y muestra la opción Cerrar Sesión.",
        ],
        "componentes": [
            ["Tasks", "AbrirSitio, Registrarse"],
            ["Interactions", "AceptarCookies, ClickJs, AbrirMenuCuenta"],
            ["Questions / Ensure", "Text.of(PANEL_CUENTA), Ensure.thatTheCurrentPage().currentUrl(), Ensure.that(ENLACE_CERRAR_SESION).isDisplayed()"],
            ["User interfaces", "MiCuentaPage (#reg_username, #first_name, #last_name, #reg_email, #reg_password, #reg_password2, #privacy_policy_reg, button[name='register'], #user-icon-wrap, #header-account-menu)"],
            ["Models / utils", "Usuario, GeneradorDatos"],
            ["Step definitions", "CuentaStepDefinitions"],
        ],
        "tag": "@Registro",
        "capturas": [(2, "Panel de la cuenta después del registro"), (4, "Menú de la cuenta con la opción Cerrar Sesión")],
    },
    {
        "hoja": "Esc2 - Modificar datos", "clave": "modificar", "id": "ESC-02", "feature": "modificar_datos.feature",
        "nombre": "Modificar los datos personales desde la pestaña Datos de Mi cuenta",
        "modulo": "Mi cuenta",
        "historia": "Como cliente registrada de Bon-bonite quiero actualizar los datos personales de mi cuenta para mantener mi información al día.",
        "prioridad": "Alta", "tipo": "Funcional E2E – camino feliz (Esquema del escenario con Ejemplos)",
        "precondiciones": "• Usuario registrado y autenticado. El antecedente lo crea por la UI, cierra sesión y vuelve a iniciar sesión con cédula y contraseña "
                          "(así también se prueban el login y el logout).",
        "datos": "Tabla de Ejemplos:\n• Nombres: Sofia Andrea\n• Apellidos: Ramirez Gomez\n• Teléfono: 3001234567",
        "esperado": "Se muestra \"Datos personales actualizados correctamente\" y, al recargar la página, los nombres, apellidos y teléfono quedan con los valores nuevos.",
        "esperados_pasos": [
            "Usuario registrado, sesión cerrada y nueva sesión iniciada con cédula y contraseña; se muestra el saludo de la cuenta.",
            "Al pulsar Actualizar Información los datos se vuelven editables; se diligencian y se pulsa Guardar.",
            "Se muestra el mensaje de confirmación.",
            "Tras recargar la página los valores guardados coinciden con los ingresados.",
        ],
        "componentes": [
            ["Tasks", "AbrirSitio, Registrarse, CerrarSesion, IniciarSesion, ModificarDatos"],
            ["Interactions", "AceptarCookies, ClickJs, AbrirMenuCuenta"],
            ["Questions / Ensure", "Text.of(MENSAJE_EXITO), Text.of(VALOR_GUARDADO.of(campo))"],
            ["User interfaces", "EditarCuentaPage (#profile-update-form .update-info-btn, input[name='first_name' | 'last_name' | 'billing_phone'], .save-info-btn, #profile-message, .profile-field[data-field])"],
            ["Models / utils", "Usuario, GeneradorDatos"],
            ["Step definitions", "CuentaStepDefinitions, ModificarDatosStepDefinitions"],
        ],
        "tag": "@ModificarDatos",
        "capturas": [(3, "Mensaje de confirmación tras guardar"), (4, "Datos persistidos después de recargar la página")],
    },
    {
        "hoja": "Esc3 - Compra", "clave": "compra", "id": "ESC-03", "feature": "compra.feature",
        "nombre": "Compra de un producto de la categoría Zapatos hasta la confirmación del pedido",
        "modulo": "Zapatos · Carrito · Checkout",
        "historia": "Como cliente registrada de Bon-bonite quiero agregar productos al carrito y llegar al pago para comprar los productos que necesito.",
        "prioridad": "Alta", "tipo": "Funcional E2E – camino feliz",
        "precondiciones": "• Usuario registrado y autenticado (antecedente: registro, logout y login).\n• Carrito vacío (usuario nuevo).\n"
                          "• Al menos un producto con talla en inventario en la categoría Zapatos.",
        "datos": "• Categoría: Zapatos (desde el menú principal)\n• Producto y talla: el primero con inventario (dinámico)\n"
                 "• Facturación/envío: Género Mujer · Teléfono 300XXXXXXX · Departamento Antioquia · Ciudad Medellín · Dirección aleatoria\n"
                 "• Autorización de tratamiento de datos: marcada",
        "esperado": "El carrito contiene solo el producto elegido, con cantidad 1 y su precio. En el checkout, el resumen muestra el producto, "
                    "\"Cantidad: 1\", el total igual al precio y el botón \"Registrar Orden\" disponible.",
        "esperados_pasos": [
            "Usuario autenticado en el sitio.",
            "Se abre la categoría Zapatos y se elige el primer producto con una talla en inventario.",
            "Se selecciona la talla, se pulsa Añadir al carrito y se abre el carrito.",
            "El carrito tiene 1 ítem con el nombre, la cantidad 1 y el precio del producto.",
            "Se avanza por los pasos del checkout y se diligencian los datos de facturación y envío (incluye listas select2 de departamento y ciudad).",
            "El resumen muestra el producto, la cantidad 1, el total correcto y el botón Registrar Orden (no se pulsa).",
        ],
        "componentes": [
            ["Tasks", "AbrirSitio, Registrarse, CerrarSesion, IniciarSesion, SeleccionarProducto, AgregarAlCarrito, IrAlPago"],
            ["Interactions", "AceptarCookies, ClickJs, AbrirMenuCuenta, SeleccionarTalla, SeleccionarOpcionSelect2"],
            ["Questions / Ensure", "TallaDisponible (lee data-product_variations), Text.of / Value.of de carrito y resumen"],
            ["User interfaces", "CategoriaPage, ProductoPage, CarritoPage, CheckoutPage"],
            ["Models / utils", "Usuario, DatosEnvio, ProductoSeleccionado, GeneradorDatos, Constantes"],
            ["Step definitions", "CuentaStepDefinitions, CompraStepDefinitions"],
        ],
        "tag": "@Compra",
        "capturas": [(2, "Producto seleccionado"), (4, "Carrito con el producto"), (6, "Checkout diligenciado y resumen del pedido")],
    },
]

A = [24, 46, 50, 40, 40, 12, 12]
for esc in ESCENARIOS:
    r = por_feature[esc["clave"]]
    ws = wb.create_sheet(esc["hoja"])
    ancho(ws, A)
    f = titulo(ws, f"b) {esc['id']} – {esc['nombre']}", f"Solución del escenario automatizado · Etiqueta {esc['tag']}", 7)
    f = seccion(ws, f, "Ficha del caso de prueba", 7)
    # La ficha empieza en la fila 5: Estado queda en B12 y Duración en B13 (referenciadas desde la Portada)
    assert f == 5
    f = ficha(ws, f, [
        ("ID", esc["id"]),
        ("Nombre", esc["nombre"]),
        ("Módulo", esc["modulo"]),
        ("Historia de usuario", esc["historia"]),
        ("Prioridad / Tipo", f"{esc['prioridad']} · {esc['tipo']}"),
        ("Precondiciones", esc["precondiciones"]),
        ("Datos de prueba", esc["datos"]),
        ("Estado", ESTADO.get(r["resultado"], r["resultado"])),
        ("Duración (s)", round(r["duracion"], 1)),
        ("Fecha / hora de ejecución", r["inicio"] + " (America/Bogota)"),
        ("Resultado esperado", esc["esperado"]),
    ], 7, A)
    est = ws["B12"]
    est.fill = fill_ok if est.value == "Pasó" else fill_ko
    est.font = f_bold
    ws["B13"].number_format = "0.0"
    ws["B13"].alignment = Alignment(horizontal="left", vertical="top")

    f = seccion(ws, f, "Escenario en Gherkin (src/test/resources/features/" + esc["feature"] + ")", 7)
    f = bloque_codigo(ws, f, leer_feature(esc["feature"]), 7)

    f = seccion(ws, f, "Pasos ejecutados y resultados", 7)
    filas = []
    for i, p in enumerate(r["pasos"]):
        acciones = "\n".join("• " + a for a in p["acciones"]) or "—"
        esperado = esc["esperados_pasos"][i] if i < len(esc["esperados_pasos"]) else ""
        obtenido = "Cumple: el paso se ejecutó y sus validaciones fueron exitosas." if p["resultado"] == "SUCCESS" \
            else "No cumple: ver reporte Serenity."
        filas.append([i + 1, p["descripcion"], acciones, esperado, obtenido,
                      ESTADO.get(p["resultado"], p["resultado"]), round(p["duracion"], 1)])
    f = tabla(ws, f, ["#", "Paso Gherkin", "Acciones automatizadas (Screenplay)", "Resultado esperado",
                      "Resultado obtenido", "Estado", "Duración (s)"], filas, A, estado_col=5)

    f = seccion(ws, f, "Componentes Screenplay utilizados", 7)
    f = tabla(ws, f, ["Capa", "Clases"], esc["componentes"], A, merges={1: 6})

    f = seccion(ws, f, "Ejecución", 7)
    f = bloque_codigo(ws, f, f"cd bonbonite-automation\nmvn clean verify -Dtags=\"{esc['tag']}\"\n"
                             f"# Reporte: target/site/serenity/index.html", 7)

    f = seccion(ws, f, "Evidencias de la ejecución", 7)
    for n, (idx_paso, desc) in enumerate(esc["capturas"], start=1):
        captura = r["pasos"][idx_paso - 1]["captura"]
        if not captura:
            continue
        nombre = f"{esc['id'].lower()}_{n}.png"
        ruta = copiar_evidencia(captura, nombre)
        ws.cell(f, 1, f"Evidencia {n}: {desc} (evidencias\\{nombre})").font = f_bold
        f += 1
        h_px = insertar_imagen(ws, ruta, f"A{f}")
        f += int(h_px / 20) + 2

# =====================================================================
# Hallazgos
# =====================================================================
ws = wb.create_sheet("Hallazgos")
A = [10, 30, 60, 16, 16, 40]
ancho(ws, A)
f = titulo(ws, "Hallazgos de la ejecución", "Observaciones y posibles defectos detectados al diseñar y ejecutar los escenarios", 6)
f = tabla(ws, f, ["ID", "Módulo", "Descripción", "Severidad", "Tipo", "Recomendación"], [
    ["H-01", "Mi cuenta › Datos", "Al pulsar \"Actualizar Información\" y guardar sin tocar el campo Género, el sistema guarda \"Masculino\" aunque el "
                                  "usuario nunca eligió un género (el select toma su primera opción por defecto). Ver evidencia esc-02_1.png.",
     "Media", "Posible defecto", "Agregar una opción vacía \"Selecciona…\" y no guardar el género si el usuario no lo elige."],
    ["H-02", "Checkout", "Aunque el usuario registró nombres, apellidos y cédula, el checkout pide de nuevo género, teléfono, departamento, ciudad "
                         "y dirección; el género del checkout por defecto es \"Hombre\".",
     "Baja", "Usabilidad", "Prellenar desde el perfil y guardar la dirección en Mi cuenta tras la primera compra."],
    ["H-03", "Catálogo", "Hay productos visibles en la categoría con todas sus tallas agotadas (p. ej. Zueco en cuero tipo bambi print); "
                         "el cliente solo lo descubre en el detalle del producto.",
     "Baja", "Usabilidad / negocio", "Marcar \"Agotado\" en la tarjeta del listado u ocultar productos sin inventario."],
    ["H-04", "Ambiente", "No existe un ambiente de pruebas; la automatización crea usuarios reales en producción y no puede confirmar pedidos.",
     "Media", "Proceso", "Habilitar un ambiente de staging con pasarela en modo sandbox para cubrir el pago de punta a punta."],
    ["H-05", "Home", "La página de inicio tarda más de un minuto en terminar de cargar (evento load); en una ejecución superó los "
                     "3 minutos y el navegador agotó su tiempo de espera. El HTML llega en menos de 1 s, así que la demora está en los "
                     "recursos que se cargan después (la página incluye contenido multimedia y un feed de Instagram).",
     "Media", "Rendimiento", "Medir la página con Lighthouse o PageSpeed para identificar los recursos lentos y aplicarles carga diferida (lazy loading)."],
], A)

# Formato de impresión
for hoja in wb.worksheets:
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(SALIDA)
print("Generado:", SALIDA)
